"""
Analyze Word template to extract placeholders
"""
from docx import Document
import re

def analyze_template(docx_path: str):
    """Extract all placeholders from Word template"""
    doc = Document(docx_path)
    placeholders = set()

    print("=== TEMPLATE ANALYSIS ===\n")

    # Analyze paragraphs
    print("Paragraphs with potential placeholders:")
    for i, para in enumerate(doc.paragraphs):
        text = para.text
        if '{{' in text or '<<' in text or '[' in text:
            print(f"\nPara {i}: {text[:100]}...")
            # Find placeholders with different formats
            found = re.findall(r'\{\{([^}]+)\}\}', text)
            found += re.findall(r'<<([^>]+)>>', text)
            found += re.findall(r'\[([^\]]+)\]', text)
            placeholders.update(found)

    # Analyze tables
    print("\n\nTables with potential placeholders:")
    for table_idx, table in enumerate(doc.tables):
        for row_idx, row in enumerate(table.rows):
            for cell_idx, cell in enumerate(row.cells):
                text = cell.text
                if '{{' in text or '<<' in text or '[' in text:
                    print(f"\nTable {table_idx}, Row {row_idx}, Cell {cell_idx}: {text[:100]}")
                    found = re.findall(r'\{\{([^}]+)\}\}', text)
                    found += re.findall(r'<<([^>]+)>>', text)
                    found += re.findall(r'\[([^\]]+)\]', text)
                    placeholders.update(found)

    print("\n\n=== UNIQUE PLACEHOLDERS FOUND ===")
    for placeholder in sorted(placeholders):
        print(f"  - {placeholder}")

    print(f"\n\nTotal unique placeholders: {len(placeholders)}")

    return placeholders

if __name__ == "__main__":
    template_path = r"C:\Users\Usuario\Finkargo_Automation_Hub\FK COL - GM - Activos.docx"
    placeholders = analyze_template(template_path)
