"""
Commission Excel Export Service

Service for generating Excel reports of broker commissions.
Creates multi-sheet Excel workbooks with:
- Detalle: Individual commission records
- Resumen por Broker: Aggregated totals by broker
- Información: Report metadata and period summary
"""

import io
import logging
from typing import List, Dict, Any
from datetime import datetime
from decimal import Decimal
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

# Spanish month names
MESES = [
    'Enero', 'Febrero', 'Marzo', 'Abril', 'Mayo', 'Junio',
    'Julio', 'Agosto', 'Septiembre', 'Octubre', 'Noviembre', 'Diciembre'
]


class ComisionExcelService:
    """
    Service for generating commission Excel reports.

    Creates professional Excel files with:
    - Formatted headers with Finkargo colors
    - Currency formatting for monetary columns
    - Percentage formatting for commission rates
    - Multiple sheets for detail and summary views
    """

    # Finkargo brand colors
    HEADER_FONT = Font(name='Arial', size=11, bold=True, color='FFFFFF')
    HEADER_FILL = PatternFill(start_color='0C147B', end_color='0C147B', fill_type='solid')
    HEADER_ALIGNMENT = Alignment(horizontal='center', vertical='center', wrap_text=True)

    DATA_FONT = Font(name='Arial', size=10)
    DATA_ALIGNMENT = Alignment(horizontal='left', vertical='center')
    NUMBER_ALIGNMENT = Alignment(horizontal='right', vertical='center')

    TOTAL_FONT = Font(name='Arial', size=11, bold=True)
    TOTAL_FILL = PatternFill(start_color='E5E7EB', end_color='E5E7EB', fill_type='solid')

    INFO_LABEL_FONT = Font(name='Arial', size=10, bold=True)
    INFO_VALUE_FONT = Font(name='Arial', size=10)

    BORDER_THIN = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    def __init__(self):
        """Initialize the Excel service."""
        pass

    def generate_comisiones_excel(
        self,
        comisiones: List[Dict[str, Any]],
        resumen: List[Dict[str, Any]],
        periodo_mes: int,
        periodo_anio: int,
        tipo_cambio: float
    ) -> io.BytesIO:
        """
        Generate Excel report for broker commissions.

        Args:
            comisiones: List of commission records (dict format from repository)
            resumen: List of broker summary records (aggregated by broker)
            periodo_mes: Period month (1-12)
            periodo_anio: Period year
            tipo_cambio: Exchange rate used

        Returns:
            BytesIO: Excel file content as bytes
        """
        try:
            wb = Workbook()

            # Sheet 1: Detalle (Commission Details)
            ws_detail = wb.active
            ws_detail.title = "Detalle"
            self._create_detail_sheet(ws_detail, comisiones)

            # Sheet 2: Resumen por Broker
            ws_resumen = wb.create_sheet(title="Resumen por Broker")
            self._create_resumen_sheet(ws_resumen, resumen)

            # Sheet 3: Información
            ws_info = wb.create_sheet(title="Información")
            self._create_info_sheet(
                ws_info,
                periodo_mes,
                periodo_anio,
                tipo_cambio,
                len(comisiones),
                resumen
            )

            # Save to BytesIO
            excel_buffer = io.BytesIO()
            wb.save(excel_buffer)
            excel_buffer.seek(0)

            logger.info(
                f"Generated commission Excel: {len(comisiones)} records, "
                f"{len(resumen)} brokers for {MESES[periodo_mes-1]} {periodo_anio}"
            )

            return excel_buffer

        except Exception as e:
            logger.error(f"Error generating commission Excel: {str(e)}")
            raise

    def _create_detail_sheet(
        self,
        ws,
        comisiones: List[Dict[str, Any]]
    ) -> None:
        """Create the Detalle (detail) sheet with individual commission records."""

        # Headers
        headers = [
            "Cliente",
            "NIT",
            "Broker",
            "Tipo",
            "Línea Crédito",
            "% Comisión Cliente",
            "Monto Comisión",
            "% Broker",
            "Operaciones Mes",
            "Monto USD",
            "T/C",
            "Monto MXN",
            "% Pago Cliente",
            "Estado"
        ]

        # Add headers
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.value = header
            cell.font = self.HEADER_FONT
            cell.fill = self.HEADER_FILL
            cell.alignment = self.HEADER_ALIGNMENT
            cell.border = self.BORDER_THIN

        # Add data rows
        for row_idx, comision in enumerate(comisiones, start=2):
            # Get broker name from joined data or use ID
            broker_info = comision.get('brokers', {})
            broker_nombre = broker_info.get('nombre', 'Unknown') if broker_info else comision.get('broker_id', 'Unknown')

            row_data = [
                comision.get('cliente_nombre', ''),
                comision.get('cliente_nit', ''),
                broker_nombre,
                self._format_tipo_comision(comision.get('tipo_comision', '')),
                comision.get('linea_credito'),
                comision.get('porcentaje_comision_cliente'),
                comision.get('monto_comision_cliente'),
                comision.get('porcentaje_broker'),
                comision.get('operaciones_mes'),
                comision.get('monto_broker_usd'),
                comision.get('tipo_cambio'),
                comision.get('monto_broker_mxn'),
                comision.get('cliente_pago_pct'),
                self._format_estado(comision.get('estado', ''))
            ]

            for col_idx, value in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.value = value if value is not None else ''
                cell.font = self.DATA_FONT
                cell.border = self.BORDER_THIN

                # Format specific columns
                if col_idx in [5, 7, 9, 10, 12]:  # Currency columns
                    cell.alignment = self.NUMBER_ALIGNMENT
                    if value is not None:
                        cell.number_format = '$#,##0.00'
                elif col_idx in [6, 8, 13]:  # Percentage columns
                    cell.alignment = self.NUMBER_ALIGNMENT
                    if value is not None:
                        cell.number_format = '0.00%'
                        cell.value = float(value) / 100 if value else 0
                elif col_idx == 11:  # Exchange rate
                    cell.alignment = self.NUMBER_ALIGNMENT
                    if value is not None:
                        cell.number_format = '0.0000'
                else:
                    cell.alignment = self.DATA_ALIGNMENT

        # Adjust column widths
        column_widths = [25, 15, 25, 12, 15, 15, 15, 12, 15, 15, 10, 15, 12, 12]
        for col_idx, width in enumerate(column_widths, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = width

    def _create_resumen_sheet(
        self,
        ws,
        resumen: List[Dict[str, Any]]
    ) -> None:
        """Create the Resumen por Broker sheet with aggregated totals."""

        # Headers
        headers = [
            "Broker",
            "Total Apertura USD",
            "Total Apertura MXN",
            "Total Operativa USD",
            "Total Operativa MXN",
            "Total USD",
            "Total MXN",
            "# Comisiones"
        ]

        # Add headers
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.value = header
            cell.font = self.HEADER_FONT
            cell.fill = self.HEADER_FILL
            cell.alignment = self.HEADER_ALIGNMENT
            cell.border = self.BORDER_THIN

        # Add data rows
        for row_idx, broker_summary in enumerate(resumen, start=2):
            row_data = [
                broker_summary.get('broker_nombre', 'Unknown'),
                broker_summary.get('total_apertura_usd', 0),
                broker_summary.get('total_apertura_mxn', 0),
                broker_summary.get('total_operativa_usd', 0),
                broker_summary.get('total_operativa_mxn', 0),
                broker_summary.get('total_usd', 0),
                broker_summary.get('total_mxn', 0),
                broker_summary.get('num_comisiones', 0)
            ]

            for col_idx, value in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.value = value
                cell.font = self.DATA_FONT
                cell.border = self.BORDER_THIN

                # Format columns
                if col_idx == 1:  # Broker name
                    cell.alignment = self.DATA_ALIGNMENT
                elif col_idx == 8:  # Count
                    cell.alignment = self.NUMBER_ALIGNMENT
                else:  # Currency columns
                    cell.alignment = self.NUMBER_ALIGNMENT
                    cell.number_format = '$#,##0.00'

        # Add totals row
        if resumen:
            total_row = len(resumen) + 2
            total_data = [
                "TOTAL",
                sum(r.get('total_apertura_usd', 0) for r in resumen),
                sum(r.get('total_apertura_mxn', 0) for r in resumen),
                sum(r.get('total_operativa_usd', 0) for r in resumen),
                sum(r.get('total_operativa_mxn', 0) for r in resumen),
                sum(r.get('total_usd', 0) for r in resumen),
                sum(r.get('total_mxn', 0) for r in resumen),
                sum(r.get('num_comisiones', 0) for r in resumen)
            ]

            for col_idx, value in enumerate(total_data, start=1):
                cell = ws.cell(row=total_row, column=col_idx)
                cell.value = value
                cell.font = self.TOTAL_FONT
                cell.fill = self.TOTAL_FILL
                cell.border = self.BORDER_THIN

                if col_idx == 1:
                    cell.alignment = Alignment(horizontal='right', vertical='center')
                elif col_idx == 8:
                    cell.alignment = self.NUMBER_ALIGNMENT
                else:
                    cell.alignment = self.NUMBER_ALIGNMENT
                    cell.number_format = '$#,##0.00'

        # Adjust column widths
        column_widths = [30, 18, 18, 18, 18, 18, 18, 12]
        for col_idx, width in enumerate(column_widths, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = width

    def _create_info_sheet(
        self,
        ws,
        periodo_mes: int,
        periodo_anio: int,
        tipo_cambio: float,
        num_comisiones: int,
        resumen: List[Dict[str, Any]]
    ) -> None:
        """Create the Información sheet with report metadata."""

        # Title
        ws['A1'] = "Reporte de Comisiones - Brokers"
        ws['A1'].font = Font(name='Arial', size=16, bold=True, color='0C147B')
        ws.merge_cells('A1:B1')

        # Calculate totals
        total_usd = sum(r.get('total_usd', 0) for r in resumen)
        total_mxn = sum(r.get('total_mxn', 0) for r in resumen)

        # Info rows
        info_data = [
            ("Período:", f"{MESES[periodo_mes-1]} {periodo_anio}"),
            ("Tipo de Cambio:", f"{tipo_cambio:.4f} MXN/USD"),
            ("Fecha TC:", datetime.now().strftime('%Y-%m-%d')),
            ("Generado:", datetime.now().strftime('%Y-%m-%d %H:%M:%S')),
            ("", ""),
            ("Total Comisiones:", str(num_comisiones)),
            ("Total Brokers:", str(len(resumen))),
            ("Total USD:", f"${total_usd:,.2f}"),
            ("Total MXN:", f"${total_mxn:,.2f}"),
        ]

        for row_idx, (label, value) in enumerate(info_data, start=3):
            label_cell = ws.cell(row=row_idx, column=1)
            label_cell.value = label
            label_cell.font = self.INFO_LABEL_FONT
            label_cell.alignment = Alignment(horizontal='right', vertical='center')

            value_cell = ws.cell(row=row_idx, column=2)
            value_cell.value = value
            value_cell.font = self.INFO_VALUE_FONT
            value_cell.alignment = Alignment(horizontal='left', vertical='center')

        # Adjust column widths
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 30

    def _format_tipo_comision(self, tipo: str) -> str:
        """Format commission type for display."""
        tipo_map = {
            'apertura': 'Apertura',
            'operativa': 'Operativa'
        }
        return tipo_map.get(tipo, tipo)

    def _format_estado(self, estado: str) -> str:
        """Format commission status for display."""
        estado_map = {
            'calculado': 'Calculado',
            'aprobado': 'Aprobado',
            'pagado': 'Pagado'
        }
        return estado_map.get(estado, estado)


# Singleton instance
_comision_excel_service_instance = None


def get_comision_excel_service() -> ComisionExcelService:
    """
    Get the singleton instance of the ComisionExcelService.

    Returns:
        ComisionExcelService: Service instance
    """
    global _comision_excel_service_instance
    if _comision_excel_service_instance is None:
        _comision_excel_service_instance = ComisionExcelService()
    return _comision_excel_service_instance
