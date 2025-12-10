"""
Contract Extractor Service
Extracts broker contract data from DOCX and PDF files using regex patterns.
Supports multiple contract versions (Marzo 2024 and Current).
"""
import re
import io
import logging
from datetime import datetime
from typing import Optional

from docx import Document
import fitz  # PyMuPDF

from src.interface.alianzas_dtos import BrokerContractData, ExtractionMethod

logger = logging.getLogger(__name__)


class ContractExtractorService:
    """
    Service for extracting broker contract data from uploaded documents.
    Uses regex patterns to identify key contract fields.
    """

    # Total number of extractable fields for confidence calculation
    TOTAL_FIELDS = 8

    # Regex patterns for contract data extraction
    # Patterns are case-insensitive and support both contract versions

    # Opening commission patterns
    OPENING_COMMISSION_PATTERNS = [
        # Marzo 2024 version
        r'[Bb]ono\s+equivalente\s+al\s+(\d+(?:[.,]\d+)?)\s*%',
        # Current version
        r'(\d+(?:[.,]\d+)?)\s*%\s+de\s+la\s+comisi[oó]n\s+de\s+apertura',
        # Generic patterns
        r'comisi[oó]n\s+(?:de\s+)?apertura[:\s]+(\d+(?:[.,]\d+)?)\s*%',
        r'apertura[:\s]+(\d+(?:[.,]\d+)?)\s*%',
    ]

    # Operational commission patterns
    OPERATIONAL_COMMISSION_PATTERNS = [
        # Marzo 2024 version
        r'(\d+(?:[.,]\d+)?)\s*%\s*(?:.*?)operativa',
        # Current version
        r'(\d+(?:[.,]\d+)?)\s*%\s*(?:.*?)por\s+operaci[oó]n',
        # Generic patterns
        r'comisi[oó]n\s+operativa[:\s]+(\d+(?:[.,]\d+)?)\s*%',
        r'operativa[:\s]+(\d+(?:[.,]\d+)?)\s*%',
    ]

    # Contract date patterns
    DATE_PATTERNS = [
        r'(?:firmado|fecha)[:\s]*(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})',
        r'(?:fecha\s+de\s+contrato|fecha\s+del\s+contrato)[:\s]*(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})',
        r'celebrado\s+(?:el\s+)?(?:d[ií]a\s+)?(\d{1,2})\s+de\s+(\w+)\s+de[l]?\s+(\d{4})',
    ]

    # Duration patterns
    DURATION_PATTERNS = [
        r'vigencia\s+(?:de\s+)?(\d+)\s*meses',
        r'plazo\s+(?:de\s+)?(\d+)\s*meses',
        r'duraci[oó]n\s+(?:de\s+)?(\d+)\s*meses',
        r'(\d+)\s*meses\s+de\s+vigencia',
    ]

    # RFC patterns (Mexican tax ID)
    RFC_PATTERNS = [
        r'R\.?F\.?C\.?[:\s]*([A-Z]{3,4}\d{6}[A-Z0-9]{3})',
        r'RFC[:\s]+([A-Z]{3,4}\d{6}[A-Z0-9]{3})',
        r'Registro\s+Federal[:\s]*([A-Z]{3,4}\d{6}[A-Z0-9]{3})',
    ]

    # CLABE bank account patterns (18 digits)
    CLABE_PATTERNS = [
        r'CLABE[:\s]*(\d{18})',
        r'cuenta\s+CLABE[:\s]*(\d{18})',
        r'n[uú]mero\s+de\s+cuenta[:\s]*(\d{18})',
    ]

    # Bank name patterns
    BANK_PATTERNS = [
        r'(?:banco|instituci[oó]n\s+bancaria)[:\s]*([A-Za-z\s]+?)(?:\n|,|\.|\s{2,})',
        r'(?:banco|instituci[oó]n)[:\s]+([A-Z][A-Za-z\s]*)',
    ]

    # Broker name patterns
    BROKER_NAME_PATTERNS = [
        r'(?:nombre\s+del\s+broker|broker)[:\s]*([A-Za-z\s]+?)(?:\n|,|\.|\s{2,})',
        r'(?:raz[oó]n\s+social)[:\s]*([A-Za-z\s\.,]+?)(?:\n|\s{2,})',
    ]

    # Spanish month names for date parsing
    SPANISH_MONTHS = {
        'enero': 1, 'febrero': 2, 'marzo': 3, 'abril': 4,
        'mayo': 5, 'junio': 6, 'julio': 7, 'agosto': 8,
        'septiembre': 9, 'octubre': 10, 'noviembre': 11, 'diciembre': 12
    }

    def extract_from_docx(self, file_bytes: bytes) -> BrokerContractData:
        """
        Extract contract data from a DOCX file.

        Args:
            file_bytes: The DOCX file content as bytes

        Returns:
            BrokerContractData with extracted fields
        """
        logger.info("Extracting contract data from DOCX file")

        try:
            # Load document from bytes
            doc = Document(io.BytesIO(file_bytes))

            # Extract text from all paragraphs
            text_parts = []
            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text)

            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for para in cell.paragraphs:
                            if para.text.strip():
                                text_parts.append(para.text)

            full_text = '\n'.join(text_parts)
            logger.debug(f"Extracted {len(full_text)} characters from DOCX")

            return self._extract_with_patterns(full_text)

        except Exception as e:
            logger.error(f"Error extracting from DOCX: {str(e)}")
            return BrokerContractData(
                extraction_method=ExtractionMethod.STANDARD,
                extraction_confidence=0.0,
                raw_text_preview=f"Error: {str(e)}"
            )

    def extract_from_pdf(self, file_bytes: bytes) -> BrokerContractData:
        """
        Extract contract data from a PDF file with selectable text.

        Args:
            file_bytes: The PDF file content as bytes

        Returns:
            BrokerContractData with extracted fields
        """
        logger.info("Extracting contract data from PDF file")

        try:
            # Open PDF from bytes
            pdf_document = fitz.open(stream=file_bytes, filetype="pdf")

            # Extract text from all pages
            text_parts = []
            for page_num in range(len(pdf_document)):
                page = pdf_document[page_num]
                page_text = page.get_text()
                if page_text.strip():
                    text_parts.append(page_text)

            pdf_document.close()

            full_text = '\n'.join(text_parts)
            logger.debug(f"Extracted {len(full_text)} characters from PDF")

            # Check if we got meaningful text
            if len(full_text.strip()) < 50:
                logger.warning("PDF appears to be scanned/image-based (little selectable text)")
                return BrokerContractData(
                    extraction_method=ExtractionMethod.STANDARD,
                    extraction_confidence=0.0,
                    raw_text_preview="PDF appears to be scanned. Consider using AI extraction for scanned documents."
                )

            return self._extract_with_patterns(full_text)

        except Exception as e:
            logger.error(f"Error extracting from PDF: {str(e)}")
            return BrokerContractData(
                extraction_method=ExtractionMethod.STANDARD,
                extraction_confidence=0.0,
                raw_text_preview=f"Error: {str(e)}"
            )

    def _extract_with_patterns(self, text: str) -> BrokerContractData:
        """
        Extract contract data from text using regex patterns.

        Args:
            text: The full text content of the document

        Returns:
            BrokerContractData with extracted fields
        """
        fields_found = 0

        # Extract opening commission
        porcentaje_apertura = self._extract_first_match(
            text, self.OPENING_COMMISSION_PATTERNS
        )
        if porcentaje_apertura is not None:
            porcentaje_apertura = self._parse_percentage(porcentaje_apertura)
            if porcentaje_apertura is not None:
                fields_found += 1
                logger.debug(f"Found opening commission: {porcentaje_apertura}%")

        # Extract operational commission
        porcentaje_operativa = self._extract_first_match(
            text, self.OPERATIONAL_COMMISSION_PATTERNS
        )
        if porcentaje_operativa is not None:
            porcentaje_operativa = self._parse_percentage(porcentaje_operativa)
            if porcentaje_operativa is not None:
                fields_found += 1
                logger.debug(f"Found operational commission: {porcentaje_operativa}%")

        # Extract contract date
        fecha_contrato = self._extract_date(text)
        if fecha_contrato:
            fields_found += 1
            logger.debug(f"Found contract date: {fecha_contrato}")

        # Extract duration
        vigencia_meses = self._extract_first_match(text, self.DURATION_PATTERNS)
        if vigencia_meses is not None:
            try:
                vigencia_meses = int(vigencia_meses)
                fields_found += 1
                logger.debug(f"Found duration: {vigencia_meses} months")
            except ValueError:
                vigencia_meses = None

        # Extract RFC
        rfc_broker = self._extract_first_match(text, self.RFC_PATTERNS)
        if rfc_broker:
            fields_found += 1
            logger.debug(f"Found RFC: {rfc_broker}")

        # Extract CLABE
        cuenta_bancaria = self._extract_first_match(text, self.CLABE_PATTERNS)
        if cuenta_bancaria:
            fields_found += 1
            logger.debug(f"Found CLABE: {cuenta_bancaria}")

        # Extract bank name
        banco = self._extract_first_match(text, self.BANK_PATTERNS)
        if banco:
            banco = banco.strip()
            fields_found += 1
            logger.debug(f"Found bank: {banco}")

        # Extract broker name
        nombre_broker = self._extract_first_match(text, self.BROKER_NAME_PATTERNS)
        if nombre_broker:
            nombre_broker = nombre_broker.strip()
            fields_found += 1
            logger.debug(f"Found broker name: {nombre_broker}")

        # Calculate confidence
        confidence = fields_found / self.TOTAL_FIELDS

        # Prepare raw text preview (first 500 chars)
        raw_preview = text[:500].replace('\n', ' ').strip() if text else None

        logger.info(f"Extraction complete: {fields_found}/{self.TOTAL_FIELDS} fields found (confidence: {confidence:.2f})")

        return BrokerContractData(
            nombre_broker=nombre_broker,
            porcentaje_comision_apertura=porcentaje_apertura,
            porcentaje_comision_operativa=porcentaje_operativa,
            fecha_contrato=fecha_contrato,
            vigencia_meses=vigencia_meses,
            rfc_broker=rfc_broker,
            cuenta_bancaria=cuenta_bancaria,
            banco=banco,
            extraction_method=ExtractionMethod.STANDARD,
            extraction_confidence=confidence,
            raw_text_preview=raw_preview
        )

    def _extract_first_match(
        self,
        text: str,
        patterns: list[str]
    ) -> Optional[str]:
        """
        Try multiple regex patterns and return the first match.

        Args:
            text: Text to search
            patterns: List of regex patterns to try

        Returns:
            First captured group from first matching pattern, or None
        """
        for pattern in patterns:
            match = re.search(pattern, text, re.IGNORECASE | re.MULTILINE)
            if match:
                return match.group(1)
        return None

    def _parse_percentage(self, value: str) -> Optional[float]:
        """
        Parse a percentage value, handling both comma and period decimals.

        Args:
            value: String representation of percentage

        Returns:
            Float value or None if parsing fails
        """
        if value is None:
            return None
        try:
            # Replace comma with period for decimal
            normalized = value.replace(',', '.')
            return float(normalized)
        except ValueError:
            return None

    def _extract_date(self, text: str) -> Optional[str]:
        """
        Extract contract date from text and convert to ISO format.

        Args:
            text: Text to search

        Returns:
            Date in ISO format (YYYY-MM-DD) or None
        """
        # Try DD/MM/YYYY or DD-MM-YYYY patterns
        for pattern in self.DATE_PATTERNS[:2]:
            match = re.search(pattern, text, re.IGNORECASE)
            if match:
                date_str = match.group(1)
                return self._parse_date_string(date_str)

        # Try Spanish format: "día de mes de año"
        pattern = self.DATE_PATTERNS[2]
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            day = int(match.group(1))
            month_name = match.group(2).lower()
            year = int(match.group(3))

            month = self.SPANISH_MONTHS.get(month_name)
            if month:
                try:
                    return datetime(year, month, day).strftime('%Y-%m-%d')
                except ValueError:
                    pass

        return None

    def _parse_date_string(self, date_str: str) -> Optional[str]:
        """
        Parse a date string in DD/MM/YYYY or DD-MM-YYYY format.

        Args:
            date_str: Date string to parse

        Returns:
            Date in ISO format (YYYY-MM-DD) or None
        """
        # Normalize separators
        normalized = date_str.replace('-', '/')

        try:
            # Try DD/MM/YYYY
            parts = normalized.split('/')
            if len(parts) == 3:
                day = int(parts[0])
                month = int(parts[1])
                year = int(parts[2])

                # Handle 2-digit years
                if year < 100:
                    year += 2000

                return datetime(year, month, day).strftime('%Y-%m-%d')
        except (ValueError, IndexError):
            pass

        return None
